"""Merge OPA motion/force CSV files into one UTF-8-BOM CSV per data group."""

from __future__ import annotations

import csv
import math
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
GROUPS = ("group1", "group2", "group3")
PATTERN = "*_opa_motion_force.csv"

EXPECTED_COLUMNS = [
    "time_ps",
    "OPA_COM_x_A",
    "OPA_COM_y_A",
    "OPA_COM_z_A",
    "OPA_disp_x_A",
    "OPA_disp_y_A",
    "OPA_disp_z_A",
    "OPA_disp_norm_A",
    "F_OPA_x_eV_A",
    "F_OPA_y_eV_A",
    "F_OPA_z_eV_A",
    "F_OPA_norm_eV_A",
]
SOURCE_COLUMNS = ["group", "system", "source_file", "frame_index"]


def source_files(group: str) -> list[Path]:
    def natural_key(path: Path) -> list[int | str]:
        relative = str(path.relative_to(DATA / group)).casefold()
        return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", relative)]

    files = sorted(
        (DATA / group).rglob(PATTERN),
        key=natural_key,
    )
    if not files:
        raise FileNotFoundError(f"No {PATTERN} files found under {DATA / group}")
    return files


def read_and_validate(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != EXPECTED_COLUMNS:
            raise ValueError(
                f"Unexpected columns in {path}:\n"
                f"expected={EXPECTED_COLUMNS}\nactual={reader.fieldnames}"
            )
        rows = list(reader)

    if not rows:
        raise ValueError(f"No data rows in {path}")
    for row_index, row in enumerate(rows):
        for column in EXPECTED_COLUMNS:
            value = row[column]
            try:
                numeric = float(value)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Non-numeric value in {path}, row {row_index + 2}, "
                    f"column {column}: {value!r}"
                ) from exc
            if not math.isfinite(numeric):
                raise ValueError(
                    f"Non-finite value in {path}, row {row_index + 2}, "
                    f"column {column}: {value!r}"
                )
    return rows


def merge_group(group: str) -> tuple[Path, int, int]:
    files = source_files(group)
    output = DATA / f"{group}_combined_opa_motion_force.csv"
    total_rows = 0

    # utf-8-sig writes a UTF-8 BOM so Microsoft Excel detects the encoding.
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SOURCE_COLUMNS + EXPECTED_COLUMNS)
        writer.writeheader()
        for path in files:
            rows = read_and_validate(path)
            system = path.parent.name
            relative_source = path.relative_to(DATA).as_posix()
            for frame_index, row in enumerate(rows):
                writer.writerow(
                    {
                        "group": group,
                        "system": system,
                        "source_file": relative_source,
                        "frame_index": frame_index,
                        **row,
                    }
                )
            total_rows += len(rows)

    return output, len(files), total_rows


def safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "data"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        tail = f"_{suffix}"
        candidate = f"{base[:31-len(tail)]}{tail}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def create_group_workbook(group: str) -> tuple[Path, int, int]:
    files = source_files(group)
    output = DATA / f"{group}.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = f"{group} OPA motion and force trajectories"
    workbook.properties.subject = "One solvent/composition per worksheet"
    workbook.properties.creator = "SAMs analysis workflow"

    used_sheet_names: set[str] = set()
    total_rows = 0
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_index, path in enumerate(files, start=1):
        rows = read_and_validate(path)
        sheet_name = safe_sheet_name(path.parent.name, used_sheet_names)
        sheet = workbook.create_sheet(title=sheet_name)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.append(EXPECTED_COLUMNS)

        for row in rows:
            sheet.append([float(row[column]) for column in EXPECTED_COLUMNS])

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30

        for column_cells in sheet.columns:
            header = str(column_cells[0].value)
            width = min(max(len(header) + 2, 14), 23)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
            for cell in column_cells[1:]:
                cell.number_format = "0.000000"

        table = Table(
            displayName=f"Table_{group}_{sheet_index}",
            ref=f"A1:L{len(rows) + 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        total_rows += len(rows)

    workbook.save(output)
    return output, len(files), total_rows


def verify_group_workbook(path: Path, expected_sheets: int, expected_rows: int) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != expected_sheets:
        raise ValueError(
            f"{path}: expected {expected_sheets} worksheets, found {len(workbook.sheetnames)}"
        )
    actual_rows = 0
    for sheet in workbook.worksheets:
        if sheet.max_column != len(EXPECTED_COLUMNS) or sheet.max_row != 102:
            raise ValueError(
                f"{path}/{sheet.title}: expected 102x12 including header, "
                f"found {sheet.max_row}x{sheet.max_column}"
            )
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if header != EXPECTED_COLUMNS:
            raise ValueError(f"{path}/{sheet.title}: header mismatch")
        actual_rows += sheet.max_row - 1
    workbook.close()
    if actual_rows != expected_rows:
        raise ValueError(f"{path}: expected {expected_rows} rows, found {actual_rows}")


def main() -> None:
    for group in GROUPS:
        output, file_count, row_count = merge_group(group)
        print(f"{group}: {file_count} files, {row_count} rows -> {output}")
        workbook_path, sheet_count, workbook_rows = create_group_workbook(group)
        verify_group_workbook(workbook_path, sheet_count, workbook_rows)
        print(
            f"{group}: {sheet_count} worksheets, {workbook_rows} rows "
            f"-> {workbook_path}"
        )


if __name__ == "__main__":
    main()
